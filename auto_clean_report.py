"""
Data Cleaning & Reporting Automation
------------------------------------
One command: raw CSV -> cleaned CSV + cleaning log + Excel report + dashboard chart.

Usage:
    python auto_clean_report.py                 # uses train.csv
    python auto_clean_report.py mydata.csv      # any other CSV (Superstore-style columns)
"""
import sys
import warnings
warnings.filterwarnings("ignore")
import os
from datetime import datetime
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.utils import get_column_letter

INPUT = sys.argv[1] if len(sys.argv) > 1 else "train.csv"
OUT = "report_output"
os.makedirs(OUT, exist_ok=True)
LOG = []   # every cleaning action is recorded here


def log(step, issue, action, rows_affected):
    LOG.append({"Step": step, "Issue found": issue, "Action taken": action,
                "Rows affected": int(rows_affected)})
    print(f"[{step}] {issue} -> {action} ({int(rows_affected)})")


# ============================================================ 1. LOAD
df = pd.read_csv(INPUT)
raw_rows, raw_cols = df.shape
print(f"Loaded {INPUT}: {raw_rows} rows x {raw_cols} columns\n")

# ============================================================ 2. CLEAN
# 2.1 Column names: strip spaces
df.columns = [c.strip() for c in df.columns]

# 2.2 Inconsistent text: trim spaces, fix case
text_cols = df.select_dtypes(include=["object", "string"]).columns
changed = 0
for c in text_cols:
    before = df[c].copy()
    df[c] = df[c].astype("string").str.strip().str.replace(r"\s+", " ", regex=True)
    changed += int((before.fillna("") != df[c].fillna("")).sum())
log("Text", "Extra / double spaces in text columns", "Trimmed and normalised spaces", changed)

for c in ["Ship Mode", "Segment", "Region", "Category", "Sub-Category"]:
    if c in df.columns:
        before = df[c].copy()
        df[c] = df[c].str.title()
        n = int((before.fillna("") != df[c].fillna("")).sum())
        if n:
            log("Text", f"Inconsistent capitalisation in '{c}'", "Converted to Title Case", n)

# 2.3 Dates: convert text -> real dates (dd/mm/yyyy)
for c in ["Order Date", "Ship Date"]:
    if c in df.columns:
        parsed = pd.to_datetime(df[c], dayfirst=True, errors="coerce")
        bad = int(parsed.isna().sum() - df[c].isna().sum())
        df[c] = parsed
        log("Dates", f"'{c}' stored as text", "Converted to datetime"
            + (f" ({bad} unreadable set to blank)" if bad else ""), len(df))

# 2.4 Numeric columns
if "Sales" in df.columns:
    before = df["Sales"].isna().sum()
    df["Sales"] = pd.to_numeric(df["Sales"], errors="coerce")
    n = int(df["Sales"].isna().sum() - before)
    if n:
        log("Numeric", "Non-numeric values in Sales", "Converted to blank", n)

# 2.5 Duplicates
exact = int(df.drop(columns=["Row ID"], errors="ignore").duplicated().sum())
df = df.drop_duplicates(subset=[c for c in df.columns if c != "Row ID"])
log("Duplicates", "Fully duplicated rows", "Removed (kept first)", exact)

key = [c for c in ["Order ID", "Product ID"] if c in df.columns]
if len(key) == 2:
    biz = int(df.duplicated(subset=key).sum())
    df = df.drop_duplicates(subset=key, keep="first")
    log("Duplicates", "Same Order ID + Product ID repeated (duplicate order lines)",
        "Removed (kept first)", biz)

# 2.6 Missing values
for c in df.columns:
    n = int(df[c].isna().sum())
    if n == 0:
        continue
    if c == "Postal Code":
        # fill from other rows with same City + State, else 0
        mode = df.groupby(["City", "State"])[c].agg(lambda s: s.mode().iloc[0] if s.notna().any() else np.nan)
        df[c] = df.apply(lambda r: mode.get((r["City"], r["State"]), np.nan) if pd.isna(r[c]) else r[c], axis=1)
        left = int(df[c].isna().sum())
        df[c] = df[c].fillna(0)
        log("Missing", "Postal Code is blank", "Filled using same City/State; unknown set to 0", n)
    elif pd.api.types.is_numeric_dtype(df[c]):
        df[c] = df[c].fillna(df[c].median())
        log("Missing", f"'{c}' has blanks", "Filled with median", n)
    elif pd.api.types.is_datetime64_any_dtype(df[c]):
        df = df.dropna(subset=[c])
        log("Missing", f"'{c}' has blank dates", "Rows removed", n)
    else:
        df[c] = df[c].fillna(df[c].mode().iloc[0])
        log("Missing", f"'{c}' has blanks", "Filled with most common value", n)

if "Postal Code" in df.columns:
    df["Postal Code"] = df["Postal Code"].astype(int)

# 2.7 Logical check: ship date cannot be before order date
if {"Order Date", "Ship Date"} <= set(df.columns):
    bad = int((df["Ship Date"] < df["Order Date"]).sum())
    df = df[df["Ship Date"] >= df["Order Date"]]
    log("Validation", "Ship Date earlier than Order Date", "Rows removed", bad)

# 2.8 Outliers: flag only (not deleted)
q1, q3 = df["Sales"].quantile([0.25, 0.75])
iqr = q3 - q1
df["Sales Outlier"] = (df["Sales"] > q3 + 3 * iqr) | (df["Sales"] < q1 - 3 * iqr)
log("Outliers", "Extremely high/low Sales (IQR x3 rule)", "Flagged in column 'Sales Outlier' (kept)",
    df["Sales Outlier"].sum())

# 2.9 Helper columns for reporting
df["Year"] = df["Order Date"].dt.year
df["Month"] = df["Order Date"].dt.to_period("M").dt.to_timestamp()
df = df.reset_index(drop=True)
df.to_csv(f"{OUT}/cleaned_data.csv", index=False)
print(f"\nCleaned: {raw_rows} -> {len(df)} rows\n")

# ============================================================ 3. SUMMARY TABLES
total_sales = df["Sales"].sum()
kpi = pd.DataFrame({
    "Metric": ["Report generated", "Source file", "Rows before cleaning", "Rows after cleaning",
               "Rows removed", "Date range", "Total sales", "Average order line value",
               "Number of customers", "Number of orders", "Top region", "Top category"],
    "Value": [datetime.now().strftime("%d-%b-%Y %H:%M"), INPUT, raw_rows, len(df), raw_rows - len(df),
              f"{df['Order Date'].min():%d-%b-%Y} to {df['Order Date'].max():%d-%b-%Y}",
              round(total_sales, 2), round(df["Sales"].mean(), 2),
              df["Customer ID"].nunique(), df["Order ID"].nunique(),
              df.groupby("Region")["Sales"].sum().idxmax(),
              df.groupby("Category")["Sales"].sum().idxmax()]})

by_cat = df.groupby("Category")["Sales"].sum().sort_values(ascending=False).round(2).reset_index()
by_reg = df.groupby("Region")["Sales"].sum().sort_values(ascending=False).round(2).reset_index()
by_seg = df.groupby("Segment")["Sales"].sum().sort_values(ascending=False).round(2).reset_index()
monthly = df.groupby("Month")["Sales"].sum().round(2).reset_index()
monthly["Month"] = monthly["Month"].dt.strftime("%Y-%m")
top_prod = (df.groupby("Product Name")["Sales"].sum().sort_values(ascending=False)
            .head(10).round(2).reset_index())
yearly = df.groupby("Year")["Sales"].sum().round(2).reset_index()
yearly["Growth %"] = (yearly["Sales"].pct_change() * 100).round(1)
log_df = pd.DataFrame(LOG)

# ============================================================ 4. EXCEL REPORT
xlsx = f"{OUT}/Sales_Report.xlsx"
with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
    kpi.to_excel(w, sheet_name="Summary", index=False)
    log_df.to_excel(w, sheet_name="Cleaning Log", index=False)
    by_cat.to_excel(w, sheet_name="By Category", index=False)
    by_reg.to_excel(w, sheet_name="By Region", index=False)
    by_seg.to_excel(w, sheet_name="By Segment", index=False)
    yearly.to_excel(w, sheet_name="Yearly", index=False)
    monthly.to_excel(w, sheet_name="Monthly Trend", index=False)
    top_prod.to_excel(w, sheet_name="Top 10 Products", index=False)

wb = load_workbook(xlsx)
head_fill = PatternFill("solid", fgColor="1F4E79")
for ws in wb.worksheets:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(width + 4, 60)
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"

# native Excel charts
ws = wb["Monthly Trend"]
ch = LineChart(); ch.title = "Monthly Sales"; ch.height, ch.width = 8, 18
ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
ws.add_chart(ch, "D2")
ws = wb["By Category"]
bc = BarChart(); bc.title = "Sales by Category"; bc.height, bc.width = 8, 14
bc.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
bc.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
ws.add_chart(bc, "D2")
wb.save(xlsx)

# ============================================================ 5. DASHBOARD (seaborn + matplotlib)
sns.set_theme(style="whitegrid")
fig, ax = plt.subplots(2, 3, figsize=(18, 10))
fig.suptitle("Sales Summary Dashboard", fontsize=18, fontweight="bold")

sns.lineplot(data=monthly, x="Month", y="Sales", ax=ax[0, 0], marker="o", color="#1F4E79")
ax[0, 0].set_title("Monthly Sales Trend")
ax[0, 0].set_xticks(ax[0, 0].get_xticks()[::6]); ax[0, 0].tick_params(axis="x", rotation=45)

sns.barplot(data=by_cat, x="Category", y="Sales", ax=ax[0, 1], palette="Blues_r")
ax[0, 1].set_title("Sales by Category")

sns.barplot(data=by_reg, x="Region", y="Sales", ax=ax[0, 2], palette="Greens_r")
ax[0, 2].set_title("Sales by Region")

ax[1, 0].pie(by_seg["Sales"], labels=by_seg["Segment"], autopct="%1.1f%%",
             colors=sns.color_palette("pastel"))
ax[1, 0].set_title("Sales by Segment")

sns.barplot(data=top_prod.head(5), y="Product Name", x="Sales", ax=ax[1, 1], palette="Oranges_r")
ax[1, 1].set_title("Top 5 Products")
ax[1, 1].set_yticklabels([t.get_text()[:28] + "..." for t in ax[1, 1].get_yticklabels()])
ax[1, 1].set_ylabel("")

sns.barplot(data=log_df[log_df["Step"] != "Dates"].groupby("Step")["Rows affected"].sum().reset_index(),
            x="Step", y="Rows affected", ax=ax[1, 2], palette="Reds_r")
ax[1, 2].set_title("Cleaning Actions (rows affected)")
ax[1, 2].tick_params(axis="x", rotation=30)

plt.tight_layout()
plt.savefig(f"{OUT}/dashboard.png", dpi=150)
print(f"Saved in '{OUT}/': cleaned_data.csv, Sales_Report.xlsx, dashboard.png")
plt.show()
